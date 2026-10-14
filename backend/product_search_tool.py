"""
Product Search & Automated Excel Reporting Subsystem
Provides natural-language product search understanding, multi-site browser extraction
(Amazon, Flipkart, and accessible shopping sources), real price/link collection,
and automatic Excel report (.xlsx) generation.
"""

import os
import re
import time
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple
from urllib.parse import quote_plus

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from utils.logger import logger, add_log_entry

REPORTS_DIR = "reports"
os.makedirs(REPORTS_DIR, exist_ok=True)


# =====================================================================
# 1. NATURAL LANGUAGE UNDERSTANDING & INTENT PARSING
# =====================================================================

KNOWN_PRODUCT_KEYWORDS = [
    "laptop", "laptops", "notebook", "gaming laptop", "gaming laptops",
    "smartphone", "smartphones", "phone", "phones", "mobile", "mobiles",
    "headphone", "headphones", "wireless headphone", "wireless headphones",
    "earphone", "earphones", "earbuds", "airpods",
    "monitor", "monitors", "display", "displays", "gaming monitor",
    "tablet", "tablets", "ipad", "ipads",
    "camera", "cameras", "dslr",
    "keyboard", "keyboards", "mechanical keyboard",
    "mouse", "mice", "gaming mouse",
    "tv", "tvs", "television", "smart tv",
    "printer", "printers", "all-in-one printer",
    "smartwatch", "smartwatches", "watch",
    "speaker", "speakers", "bluetooth speaker",
    "refrigerator", "washing machine", "ac", "air conditioner",
    "iphone", "macbook", "samsung phone", "samsung phones"
]


def parse_product_search_intent(prompt: str) -> Optional[Dict[str, Any]]:
    """
    Analyzes natural language prompt to determine if it is a product search or comparison request.
    Extracts category, maximum price, minimum price, quantity, and task type.
    """
    if not prompt or not isinstance(prompt, str):
        return None

    p_clean = prompt.strip()
    p_lower = p_clean.lower()

    # Exclude internal enterprise demo tasks
    if any(k in p_lower for k in [
        "unpaid invoice", "process invoice", "invoice reminder", "overdue invoice",
        "office supply", "order the office supplies", "procurement", "products.xlsx",
        "handle my work", "work for today", "daily work", "today's work",
        "client meeting", "prepare for my 3 pm", "meeting brief",
        "activity report", "performance report", "monthly report"
    ]):
        return None

    # Determine Task Type: Comparison vs Search
    is_comparison = any(w in p_lower for w in ["compare", "comparison", "versus", "vs", "difference in price"])
    task_type = "product_comparison" if is_comparison else "product_search"

    # Check for search/product trigger keywords
    has_search_action = any(w in p_lower for w in [
        "search", "find", "compare", "look for", "get", "list", "show", "buy",
        "price of", "prices", "under", "below", "between", "budget", "cheap", "best"
    ])

    matched_keyword = None
    for kw in sorted(KNOWN_PRODUCT_KEYWORDS, key=len, reverse=True):
        if re.search(r"\b" + re.escape(kw) + r"\b", p_lower):
            matched_keyword = kw
            break

    # If neither an explicit action keyword nor a known product is found, not a product search
    if not has_search_action and not matched_keyword:
        return None

    # 1. Extract Quantity (e.g., "Find 10 smartphones", "10 gaming laptops", "top 5")
    quantity = 10  # default
    user_specified_qty = None
    qty_match = re.search(r"\b(?:find|get|show|search for|list)\s+(\d{1,2})\s+(?:valid\s+|best\s+|top\s+)?([a-z]+)", p_lower)
    if not qty_match:
        qty_match = re.search(r"^\s*(\d{1,2})\s+(?:valid\s+|best\s+|top\s+)?([a-z]+)", p_lower)
    if qty_match:
        try:
            cand_num = int(qty_match.group(1))
            cand_word = qty_match.group(2)
            if cand_word not in ["k", "lakh", "lac", "thousand", "hundred", "rs", "inr", "under", "below", "price", "prices"] and 1 <= cand_num <= 50:
                user_specified_qty = cand_num
                quantity = cand_num
        except Exception:
            pass

    # 2. Extract Price Range / Maximum Price
    max_price = None
    min_price = None

    # Pattern A: between X and Y (e.g., between ₹20,000 and ₹40,000)
    between_match = re.search(r"between\s*(?:rs\.?|inr|₹)?\s*([\d,]+)\s*(?:and|to|-)\s*(?:rs\.?|inr|₹)?\s*([\d,]+)", p_lower)
    if between_match:
        try:
            p1 = float(between_match.group(1).replace(",", ""))
            p2 = float(between_match.group(2).replace(",", ""))
            min_price = min(p1, p2)
            max_price = max(p1, p2)
        except Exception:
            pass

    # Pattern B: under / below / max / less than ₹XX,XXX or XXk
    if max_price is None:
        under_match = re.search(r"(?:under|below|less than|within|up to|max(?:imum)?(?: price)? of?)\s*(?:rs\.?|inr|₹)?\s*([\d,]+)(?:\s*(k|lakh|lac))?", p_lower)
        if under_match:
            try:
                raw_num = under_match.group(1).replace(",", "")
                multiplier = under_match.group(2)
                val = float(raw_num)
                if multiplier == "k":
                    val *= 1000
                elif multiplier in ("lakh", "lac"):
                    val *= 100000
                max_price = val
            except Exception:
                pass

    # 3. Extract Category / Product Query
    category = ""
    filtered = p_clean
    patterns_to_strip = [
        r"(?i)\b(search\s+for|find|compare|look\s+for|show\s+me|show|get\s+me|get|list)\b",
        r"(?i)\b(under|below|less than|between|up to|max(?:imum)?(?: price)? of?)\s*(?:rs\.?|inr|₹)?\s*[\d,]+(?:\s*(?:k|lakh|lac))?",
        r"(?i)\band\s*(?:rs\.?|inr|₹)?\s*[\d,]+(?:\s*(?:k|lakh|lac))?",
        r"(?i)(?:rs\.?|inr|₹)\s*[\d,]+(?:\s*(?:k|lakh|lac))?",
        r"(?i)\b(prices?|deals?|options?|costs?)\b"
    ]
    if user_specified_qty:
        patterns_to_strip.append(rf"\b{user_specified_qty}\b")

    for pat in patterns_to_strip:
        filtered = re.sub(pat, " ", filtered)

    # Clean whitespace and punctuation
    cleaned_category = " ".join(filtered.split()).strip().strip(",- ")

    if cleaned_category:
        category = cleaned_category
    elif matched_keyword:
        category = matched_keyword
    else:
        category = "product"

    # Normalize singular for simple plurals (e.g. "laptops" -> "laptop", "smartphones" -> "smartphone")
    # while preserving model names like "iPhone 16"
    cat_lower = category.lower()
    if cat_lower.endswith("s") and not cat_lower.endswith("ss") and not re.search(r"\d+$", cat_lower):
        # Don't strip 's' if last word is a model or brand
        words = category.split()
        if len(words) == 1 and words[0].lower() in ["laptops", "smartphones", "headphones", "monitors", "tablets", "cameras", "keyboards", "mice", "printers"]:
            category = words[0][:-1]
            if words[0].lower() == "mice":
                category = "mouse"
    query = category.strip()

    add_log_entry(f"Product search request understood: Category='{category}', Max Price={f'₹{max_price:,.0f}' if max_price else 'None'}, Quantity={quantity}, Task={task_type}", "INFO")
    add_log_entry(f"Search requirements extracted: query='{query}'", "SUCCESS")

    return {
        "task_type": task_type,
        "category": category,
        "query": query,
        "maximum_price": max_price,
        "minimum_price": min_price,
        "quantity": quantity,
        "raw_prompt": prompt
    }


# =====================================================================
# 2. BROWSER AUTOMATION & REAL PRODUCT DATA EXTRACTION
# =====================================================================

def _clean_price_val(price_str: str) -> Tuple[Optional[float], str]:
    """Parses text containing currency into numeric float and formatted ₹XX,XXX string."""
    if not price_str or price_str.strip() in ("", "N/A", "None"):
        return None, "N/A"
    clean = re.sub(r"[^\d.]", "", price_str.replace(",", ""))
    try:
        val = float(clean)
        return val, f"₹{int(val):,}"
    except ValueError:
        return None, price_str.strip()


def extract_from_amazon(driver, query: str, max_price: Optional[float] = None, limit: int = 10) -> List[Dict[str, Any]]:
    """Navigates to Amazon and extracts real product cards."""
    from selenium.webdriver.common.by import By

    url = f"https://www.amazon.in/s?k={quote_plus(query)}"
    add_log_entry(f"Opening product search website (Amazon.in)", "ACTION")
    try:
        driver.get(url)
        time.sleep(2.5)
        add_log_entry("Search page loaded (Amazon)", "SUCCESS")
    except Exception as e:
        add_log_entry(f"WARNING: Amazon search page could not be accessed: {e}", "WARNING")
        return []

    add_log_entry(f"Searching for '{query}' on Amazon", "ACTION")
    add_log_entry("Extracting product information from Amazon", "ACTION")

    items = driver.find_elements(By.CSS_SELECTOR, "div[data-component-type='s-search-result']")
    products = []
    seen_titles = set()

    for it in items:
        try:
            # 1. Product Name
            title = ""
            for sel in ["h2 span", "h2 a span", "h2", "a.a-link-normal span"]:
                t_els = it.find_elements(By.CSS_SELECTOR, sel)
                if t_els and t_els[0].text.strip():
                    title = t_els[0].text.strip()
                    break
            if not title:
                continue

            # Deduplication key
            norm_key = re.sub(r"[^a-zA-Z0-9]", "", title.lower())[:35]
            if norm_key in seen_titles:
                continue

            # 2. Price
            price_val = None
            price_formatted = "N/A"
            for sel in ["span.a-price-whole", "span.a-price span.a-offscreen", "span.a-price"]:
                p_els = it.find_elements(By.CSS_SELECTOR, sel)
                if p_els and p_els[0].text.strip():
                    price_val, price_formatted = _clean_price_val(p_els[0].text.strip())
                    break

            # Filter by max price if specified
            if max_price and price_val and price_val > max_price:
                continue

            # 3. Product Link
            link = "N/A"
            for sel in ["a.a-link-normal[href*='/dp/']", "h2 a", "a.a-link-normal.s-no-outline", "a[href*='/dp/']"]:
                l_els = it.find_elements(By.CSS_SELECTOR, sel)
                if l_els:
                    h = l_els[0].get_attribute("href")
                    if h and ("/dp/" in h or "/gp/" in h):
                        link = h.split("/ref=")[0]
                        break
            if link == "N/A":
                h2_a = it.find_elements(By.CSS_SELECTOR, "h2 a")
                if h2_a:
                    link = h2_a[0].get_attribute("href") or "N/A"

            # 4. Rating
            rating = "N/A"
            for sel in ["span[aria-label*='stars']", "i.a-icon-star-small span", "span.a-icon-alt"]:
                r_els = it.find_elements(By.CSS_SELECTOR, sel)
                if r_els:
                    txt = r_els[0].get_attribute("innerText") or r_els[0].text
                    if txt:
                        rating = txt.strip()
                        break

            # 5. Original Price / Discount
            orig_price = "N/A"
            for sel in ["span.a-price.a-text-price span.a-offscreen", "span.a-text-strike"]:
                op_els = it.find_elements(By.CSS_SELECTOR, sel)
                if op_els and op_els[0].text.strip():
                    orig_price = op_els[0].text.strip()
                    break

            seen_titles.add(norm_key)
            products.append({
                "Product Name": title,
                "Website": "Amazon",
                "Price": price_formatted,
                "Product Link": link,
                "Rating": rating,
                "Original Price": orig_price,
                "_price_num": price_val or 0
            })

            if len(products) >= limit:
                break
        except Exception:
            continue

    add_log_entry(f"Collected {len(products)} authentic product listings from Amazon", "SUCCESS")
    return products


def extract_from_flipkart(driver, query: str, max_price: Optional[float] = None, limit: int = 10) -> List[Dict[str, Any]]:
    """Navigates to Flipkart and extracts real product cards."""
    from selenium.webdriver.common.by import By

    url = f"https://www.flipkart.com/search?q={quote_plus(query)}"
    add_log_entry(f"Opening product search website (Flipkart)", "ACTION")
    try:
        driver.get(url)
        time.sleep(2.5)
        add_log_entry("Search page loaded (Flipkart)", "SUCCESS")
    except Exception as e:
        add_log_entry(f"WARNING: Flipkart search page could not be accessed: {e}", "WARNING")
        return []

    add_log_entry(f"Searching for '{query}' on Flipkart", "ACTION")
    add_log_entry("Extracting product information from Flipkart", "ACTION")

    cards = driver.find_elements(By.CSS_SELECTOR, "div[data-id], div.tUxRFH, div._75nlfW, div.slAVV4, div._1sdMkc")
    products = []
    seen_titles = set()

    for c in cards:
        try:
            # 1. Product Name
            title = ""
            for sel in ["div.KzDlHZ", "a.wjcEIp", "a.WKTcLC", "div._4rR01T", "a.s1Q9rs", "div[class*='title']"]:
                t_els = c.find_elements(By.CSS_SELECTOR, sel)
                if t_els and t_els[0].text.strip():
                    title = t_els[0].text.strip()
                    break

            if not title:
                # Parse lines from card text directly
                raw_lines = [ln.strip() for ln in c.text.split("\n") if ln.strip()]
                for ln in raw_lines:
                    if len(ln) > 10 and not any(k in ln.lower() for k in ["compare", "ratings", "reviews", "top discount", "off on exchange", "free delivery"]):
                        title = ln
                        break

            if not title:
                continue

            norm_key = re.sub(r"[^a-zA-Z0-9]", "", title.lower())[:35]
            if norm_key in seen_titles:
                continue

            # 2. Price
            price_val = None
            price_formatted = "N/A"
            for sel in ["div.Nx9q8j", "div._30jeq3", "div[class*='price']"]:
                p_els = c.find_elements(By.CSS_SELECTOR, sel)
                if p_els and ("₹" in p_els[0].text or "Rs" in p_els[0].text):
                    price_val, price_formatted = _clean_price_val(p_els[0].text.strip())
                    break

            if price_formatted == "N/A":
                # Fallback to card text rupee matches
                matches = re.findall(r"₹\s*[\d,]+", c.text)
                if matches:
                    price_val, price_formatted = _clean_price_val(matches[0])

            if max_price and price_val and price_val > max_price:
                continue

            # 3. Product Link
            link = "N/A"
            for sel in ["a[href*='/p/']", "a.CGtC58", "a.wjcEIp", "a[target='_blank']"]:
                l_els = c.find_elements(By.CSS_SELECTOR, sel)
                if l_els:
                    h = l_els[0].get_attribute("href")
                    if h and "/p/" in h:
                        link = h.split("&pid=")[0] if ("&pid=" in h and "?" in h) else h
                        break

            # 4. Rating
            rating = "N/A"
            for sel in ["div.XQDdHH", "div._3LWZlK", "span[id*='productRating']"]:
                r_els = c.find_elements(By.CSS_SELECTOR, sel)
                if r_els and r_els[0].text.strip():
                    rating = r_els[0].text.strip()
                    break

            # 5. Original Price / Discount
            orig_price = "N/A"
            disc_els = c.find_elements(By.CSS_SELECTOR, "div.yRaY8j, div._3I9_wc")
            if disc_els and disc_els[0].text.strip():
                orig_price = disc_els[0].text.strip()

            seen_titles.add(norm_key)
            products.append({
                "Product Name": title,
                "Website": "Flipkart",
                "Price": price_formatted,
                "Product Link": link,
                "Rating": rating,
                "Original Price": orig_price,
                "_price_num": price_val or 0
            })

            if len(products) >= limit:
                break
        except Exception:
            continue

    add_log_entry(f"Collected {len(products)} authentic product listings from Flipkart", "SUCCESS")
    return products


def execute_multi_site_product_search(driver, intent_data: Dict[str, Any]) -> List[Dict[str, Any]]:
    """
    Executes search across accessible shopping websites, balancing targets between Amazon and Flipkart.
    Enforces budget constraints and removes intra-site duplicates.
    """
    query = intent_data.get("query", "laptop")
    max_price = intent_data.get("maximum_price")
    target_qty = intent_data.get("quantity", 10)
    task_type = intent_data.get("task_type", "product_search")

    half_qty = max(target_qty // 2, 3)

    all_products: List[Dict[str, Any]] = []

    # 1. Search Primary Source (Amazon)
    try:
        amz_results = extract_from_amazon(driver, query=query, max_price=max_price, limit=half_qty)
        all_products.extend(amz_results)
    except Exception as e:
        add_log_entry(f"WARNING: Amazon search encountered error: {e}", "WARNING")

    # 2. Search Secondary Source (Flipkart)
    remaining_needed = target_qty - len(all_products)
    fk_limit = max(remaining_needed, half_qty)
    try:
        fk_results = extract_from_flipkart(driver, query=query, max_price=max_price, limit=fk_limit)
        all_products.extend(fk_results)
    except Exception as e:
        add_log_entry(f"WARNING: Flipkart search encountered error: {e}", "WARNING")

    # 3. Deduplication & Final Validation
    add_log_entry("Validating collected products", "ACTION")
    final_products: List[Dict[str, Any]] = []
    seen_site_title = set()

    for p in all_products:
        key = (p["Website"].lower(), re.sub(r"[^a-zA-Z0-9]", "", p["Product Name"].lower())[:30])
        if key in seen_site_title:
            continue
        seen_site_title.add(key)
        final_products.append(p)
        if len(final_products) >= target_qty:
            break

    add_log_entry(f"{len(final_products)} valid products found", "SUCCESS")
    return final_products


# =====================================================================
# 3. AUTOMATIC EXCEL REPORT GENERATION (openpyxl)
# =====================================================================

def generate_product_excel(products: List[Dict[str, Any]], query_name: str = "product_search") -> str:
    """
    Automatically compiles collected product data into a formatted Excel spreadsheet (.xlsx).
    Filename format: product_search_YYYYMMDD_HHMMSS.xlsx
    """
    add_log_entry("Generating Excel report", "ACTION")

    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"product_search_{timestamp_str}.xlsx"
    filepath = os.path.join(REPORTS_DIR, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product Search Results"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Styling Palettes
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10, color="0F172A")
    link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # Columns
    columns = [
        ("Product Name", 45),
        ("Website", 15),
        ("Price", 16),
        ("Product Link", 30),
        ("Rating", 12),
        ("Original Price", 16)
    ]

    # Write Header Row
    for col_idx, (col_name, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Write Data Rows
    for row_idx, prod in enumerate(products, start=2):
        name_cell = ws.cell(row=row_idx, column=1, value=prod.get("Product Name", "N/A"))
        site_cell = ws.cell(row=row_idx, column=2, value=prod.get("Website", "N/A"))
        price_cell = ws.cell(row=row_idx, column=3, value=prod.get("Price", "N/A"))
        link_val = prod.get("Product Link", "N/A")
        link_cell = ws.cell(row=row_idx, column=4)

        if link_val and link_val != "N/A" and link_val.startswith("http"):
            link_cell.value = f'=HYPERLINK("{link_val}", "View Product")'
            link_cell.font = link_font
        else:
            link_cell.value = "N/A"
            link_cell.font = regular_font

        rating_cell = ws.cell(row=row_idx, column=5, value=prod.get("Rating", "N/A"))
        orig_price_cell = ws.cell(row=row_idx, column=6, value=prod.get("Original Price", "N/A"))

        # Formatting
        row_fill = alt_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, cell in enumerate([name_cell, site_cell, price_cell, link_cell, rating_cell, orig_price_cell], start=1):
            if cell != link_cell:
                cell.font = regular_font
            cell.border = thin_border
            if row_fill.fill_type:
                cell.fill = row_fill
            cell.alignment = Alignment(
                horizontal="center" if col_idx in (2, 3, 4, 5, 6) else "left",
                vertical="center"
            )

        ws.row_dimensions[row_idx].height = 22

    # Column Widths
    for col_idx, (_, width) in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    wb.save(filepath)
    add_log_entry(f"Excel report generated: {filename}", "SUCCESS")
    return filepath
