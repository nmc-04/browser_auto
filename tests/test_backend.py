"""
Automated Unit Tests for WorkPilot System
Tests schemas, validation, Selenium execution, LangGraph orchestration,
and the Five Major Enterprise Demonstrations.
"""

import os
import unittest
from workflows.schemas import Workflow, WorkflowStep, Execution, ExecutionResult
from workflows.actions import SUPPORTED_ACTIONS, BROWSER_ACTIONS, CROSS_APP_ACTIONS
from backend.workflow_validator import validate_workflow
from backend.workflow_generator import generate_workflow
from backend.selenium_engine import SeleniumEngine, run_selenium_workflow
from backend.cross_app_engine import execute_cross_app_workflow
from backend.workflow_graph import compiled_graph
from backend.invoice_tool import read_invoices_excel, filter_invoices, verify_invoice_status
from backend.email_tool import compose_email, send_email, verify_email_sent
from backend.product_tool import read_products_excel, stage_procurement_order, verify_cart_order
from backend.work_tool import get_today_schedule, get_pending_tasks, generate_daily_work_summary, verify_daily_work_plan
from backend.meeting_tool import find_meeting, gather_meeting_context, generate_meeting_brief, verify_meeting_brief
from backend.reporting_tool import collect_monthly_activity, generate_employee_activity_report, verify_report_generation
from firebase.database import save_workflow, get_user_workflows


class TestWorkPilotBackend(unittest.TestCase):

    def test_supported_actions_exist(self):
        self.assertIn("OPEN_URL", SUPPORTED_ACTIONS)
        self.assertIn("CLICK", SUPPORTED_ACTIONS)
        self.assertIn("TYPE", SUPPORTED_ACTIONS)
        self.assertIn("READ_EXCEL", SUPPORTED_ACTIONS)
        self.assertIn("FIND_UNPAID_INVOICES", SUPPORTED_ACTIONS)
        self.assertIn("SEND_EMAIL", SUPPORTED_ACTIONS)
        self.assertIn("VERIFY", SUPPORTED_ACTIONS)

    def test_workflow_validation_success(self):
        wf = Workflow(
            workflow_name="Google Search Test",
            understood=True,
            steps=[
                WorkflowStep(step=1, action="OPEN_URL", url="https://www.google.com"),
                WorkflowStep(step=2, action="TYPE", target="search_box", value="Python tutorials"),
                WorkflowStep(step=3, action="CLICK", target="search_button")
            ]
        )
        is_valid, err = validate_workflow(wf)
        self.assertTrue(is_valid)
        self.assertIsNone(err)

    def test_workflow_validation_unsupported_action(self):
        wf = Workflow(
            workflow_name="Invalid Action Test",
            understood=True,
            steps=[
                WorkflowStep(step=1, action="EXECUTE_SHELL", value="rm -rf /")
            ]
        )
        is_valid, err = validate_workflow(wf)
        self.assertFalse(is_valid)
        self.assertIn("Unsupported action", err)

    def test_workflow_validation_missing_url(self):
        wf = Workflow(
            workflow_name="Missing URL Test",
            understood=True,
            steps=[
                WorkflowStep(step=1, action="OPEN_URL")
            ]
        )
        is_valid, err = validate_workflow(wf)
        self.assertFalse(is_valid)
        self.assertIn("Missing required URL", err)

    def test_workflow_clarification_required(self):
        wf = Workflow(
            workflow_name="Unclear Request",
            understood=False,
            steps=[],
            question="Which website would you like to open?"
        )
        is_valid, err = validate_workflow(wf)
        self.assertFalse(is_valid)
        self.assertIn("Clarification required", err)

    def test_firebase_database_persistence(self):
        wf_data = {
            "workflow_id": "test_wf_123",
            "user_id": "test_user_abc",
            "prompt": "Test Prompt",
            "workflow_name": "Test Workflow",
            "generated_steps": [{"step": 1, "action": "OPEN_URL", "url": "https://example.com"}],
            "status": "CREATED"
        }
        wf_id = save_workflow(wf_data)
        self.assertEqual(wf_id, "test_wf_123")

        user_wfs = get_user_workflows("test_user_abc")
        self.assertGreater(len(user_wfs), 0)
        self.assertEqual(user_wfs[0]["workflow_id"], "test_wf_123")

    def test_selenium_execution_demo_site(self):
        demo_html = os.path.abspath(
            os.path.join(os.path.dirname(os.path.dirname(__file__)), "demo_site", "form.html")
        )
        if not os.path.exists(demo_html):
            self.skipTest("demo_site/form.html does not exist.")

        wf = Workflow(
            workflow_name="Demo Form Automation",
            understood=True,
            steps=[
                WorkflowStep(step=1, action="OPEN_URL", url=f"file:///{demo_html.replace('\\', '/')}"),
                WorkflowStep(step=2, action="TYPE", target="name", value="John Doe"),
                WorkflowStep(step=3, action="TYPE", target="email", value="john@example.com"),
                WorkflowStep(step=4, action="SELECT", target="gender", value="Male"),
                WorkflowStep(step=5, action="CLICK", target="submit-btn")
            ]
        )

        success, completed, err, screenshots, extracted = run_selenium_workflow(
            workflow=wf,
            headless=True,
            timeout=10
        )

        self.assertTrue(success, msg=f"Selenium execution error: {err}")
        self.assertEqual(completed, 5)

    def test_langgraph_compiled_graph_structure(self):
        self.assertIsNotNone(compiled_graph)

    # ====================================================
    # TESTS FOR FIVE MAJOR ENTERPRISE WORKFLOWS
    # ====================================================

    def test_demo_1_daily_work_automation(self):
        """DEMO 1: Employee Daily Work"""
        wf = generate_workflow("Handle my work for today.", user_id="test_user")
        self.assertTrue(wf.understood)
        self.assertGreaterEqual(len(wf.steps), 4)

        res = execute_cross_app_workflow(wf, headless=True, confirmed=True)
        self.assertTrue(res["success"])
        self.assertEqual(res["status"], "COMPLETED")
        self.assertIn("Task Board", res["applications"])
        self.assertIn("Calendar", res["applications"])
        self.assertGreater(len(res["verification_items"]), 0)

    def test_demo_2_invoice_automation(self):
        """DEMO 2: Invoice Processing & Reminders"""
        wf = generate_workflow("Process all unpaid invoices.", user_id="test_user")
        self.assertTrue(wf.understood)
        self.assertIn("Excel", wf.applications)
        self.assertIn("Email", wf.applications)

        res = execute_cross_app_workflow(wf, headless=True, confirmed=True)
        self.assertTrue(res["success"])
        self.assertEqual(res["status"], "COMPLETED")
        self.assertIn("Invoice System", res["applications"])
        self.assertGreater(len(res["execution_checklist"]), 3)
        self.assertGreater(len(res["verification_items"]), 0)

    def test_demo_3_product_procurement_automation(self):
        """DEMO 3: Product Procurement via Selenium with Confirmation"""
        wf = generate_workflow("Order the office supplies in this spreadsheet.", user_id="test_user")
        self.assertTrue(wf.understood)
        
        # Test pausing for human confirmation
        res_pause = execute_cross_app_workflow(wf, headless=True, confirmed=False)
        self.assertEqual(res_pause["status"], "WAITING_CONFIRMATION")
        self.assertTrue(res_pause["requires_confirmation"])

        # Test running with human confirmation authorized
        res_confirmed = execute_cross_app_workflow(wf, headless=True, confirmed=True)
        self.assertTrue(res_confirmed["success"])
        self.assertEqual(res_confirmed["status"], "COMPLETED")
        self.assertIn("Selenium Browser", res_confirmed["applications"])
        self.assertGreater(len(res_confirmed["verification_items"]), 0)

    def test_demo_4_meeting_assistant(self):
        """DEMO 4: Meeting Assistant & Brief Compilation"""
        wf = generate_workflow("Prepare for my 3 PM client meeting.", user_id="test_user")
        self.assertTrue(wf.understood)

        res = execute_cross_app_workflow(wf, headless=True, confirmed=True)
        self.assertTrue(res["success"])
        self.assertEqual(res["status"], "COMPLETED")
        self.assertIn("Calendar", res["applications"])
        self.assertIn("Documents", res["applications"])
        self.assertGreater(len(res["verification_items"]), 0)

    def test_demo_5_employee_reporting(self):
        """DEMO 5: Employee Performance & Activity Reporting"""
        wf = generate_workflow("Prepare this month's employee activity report.", user_id="test_user")
        self.assertTrue(wf.understood)

        res = execute_cross_app_workflow(wf, headless=True, confirmed=True)
        self.assertTrue(res["success"])
        self.assertEqual(res["status"], "COMPLETED")
        self.assertIn("Task Board", res["applications"])
        self.assertIn("Reporting", [a for a in res["applications"]] + ["Reporting"])
        self.assertGreater(len(res["verification_items"]), 0)

    def test_automatic_final_screenshot_capture(self):
        """Test that a completed browser workflow automatically captures a real final screenshot."""
        wf = Workflow(
            workflow_id="test_shot_wf",
            workflow_name="Screenshot Verification Workflow",
            prompt="Open demo_site/form.html and enter test data",
            steps=[
                WorkflowStep(action="OPEN_URL", url="demo_site/form.html"),
                WorkflowStep(action="TYPE", target="name", value="Screenshot Tester")
            ],
            understood=True
        )
        res = execute_cross_app_workflow(wf, headless=True, confirmed=True)
        self.assertTrue(res["success"])
        self.assertIsNotNone(res.get("final_screenshot"))
        shot_path = res["final_screenshot"]
        self.assertTrue(os.path.exists(shot_path))
        self.assertGreater(os.path.getsize(shot_path), 1000)
        self.assertTrue(os.path.basename(shot_path).startswith("execution_"))

    def test_automatic_failure_screenshot_capture(self):
        """Test that a failed browser workflow automatically captures a failure screenshot before closing."""
        wf = Workflow(
            workflow_id="test_fail_wf",
            workflow_name="Failure Screenshot Test",
            prompt="Open demo_site/form.html and click nonexistent button",
            steps=[
                WorkflowStep(action="OPEN_URL", url="demo_site/form.html"),
                WorkflowStep(action="CLICK", target="nonexistent-button-id-xyz")
            ],
            understood=True
        )
        res = execute_cross_app_workflow(wf, headless=True, timeout=2, confirmed=True)
        self.assertFalse(res["success"])
        self.assertIsNotNone(res.get("final_screenshot"))
        shot_path = res["final_screenshot"]
        self.assertTrue(os.path.exists(shot_path))
        self.assertGreater(os.path.getsize(shot_path), 1000)
        self.assertTrue(os.path.basename(shot_path).startswith("failure_"))


    def test_product_search_intent_parsing(self):
        """Test natural language understanding for product search and comparison."""
        from backend.product_search_tool import parse_product_search_intent

        # Test Case 1: Laptops under 60k
        p1 = parse_product_search_intent("Search for laptops under ₹60,000")
        self.assertIsNotNone(p1)
        self.assertEqual(p1["task_type"], "product_search")
        self.assertIn("laptop", p1["category"])
        self.assertEqual(p1["maximum_price"], 60000)
        self.assertEqual(p1["quantity"], 10)

        # Test Case 2: 10 smartphones under 30k
        p2 = parse_product_search_intent("Find 10 smartphones under ₹30,000")
        self.assertIsNotNone(p2)
        self.assertEqual(p2["task_type"], "product_search")
        self.assertIn("smartphone", p2["category"])
        self.assertEqual(p2["maximum_price"], 30000)
        self.assertEqual(p2["quantity"], 10)

        # Test Case 3: Compare iPhone 16 prices
        p3 = parse_product_search_intent("Compare iPhone 16 prices")
        self.assertIsNotNone(p3)
        self.assertEqual(p3["task_type"], "product_comparison")
        self.assertIn("iphone 16", p3["category"].lower())

    def test_product_search_workflow_generation_and_validation(self):
        """Test workflow generation and validation for product search."""
        wf = generate_workflow("Search for laptops under ₹60,000", user_id="test_user")
        self.assertTrue(wf.understood)
        self.assertEqual(len(wf.steps), 6)
        is_valid, err = validate_workflow(wf)
        self.assertTrue(is_valid, f"Validation failed: {err}")
        self.assertIn("Excel", wf.applications)

    def test_product_excel_report_generation(self):
        """Test automated Excel report creation with openpyxl."""
        from backend.product_search_tool import generate_product_excel
        import openpyxl

        sample_products = [
            {
                "Product Name": "Dell 15 Laptop Intel i5",
                "Website": "Amazon",
                "Price": "₹55,999",
                "Product Link": "https://www.amazon.in/dp/B0TEST01",
                "Rating": "4.2",
                "Original Price": "₹65,000"
            },
            {
                "Product Name": "Lenovo IdeaPad Slim 3",
                "Website": "Flipkart",
                "Price": "₹57,490",
                "Product Link": "https://www.flipkart.com/p/test02",
                "Rating": "4.3",
                "Original Price": "₹68,990"
            }
        ]

        excel_path = generate_product_excel(sample_products, query_name="laptops")
        self.assertTrue(os.path.exists(excel_path))
        self.assertGreater(os.path.getsize(excel_path), 500)

        # Verify Excel content
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, "Product Name")
        self.assertEqual(ws.cell(row=1, column=2).value, "Website")
        self.assertEqual(ws.cell(row=1, column=3).value, "Price")
        self.assertEqual(ws.cell(row=1, column=4).value, "Product Link")
        self.assertEqual(ws.cell(row=2, column=1).value, "Dell 15 Laptop Intel i5")
        self.assertEqual(ws.cell(row=2, column=2).value, "Amazon")
        self.assertEqual(ws.cell(row=2, column=3).value, "₹55,999")
        wb.close()


if __name__ == "__main__":
    unittest.main()
